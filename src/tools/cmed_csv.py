# Converte a planilha da lista CMED (PMC) em CSV: python3 cmed_csv.py lista.xlsx saida.csv
import openpyxl, csv, sys
wb = openpyxl.load_workbook(sys.argv[1], read_only=True)
ws = wb['Lista PMC'] if 'Lista PMC' in wb.sheetnames else wb.worksheets[-1]
w = csv.writer(open(sys.argv[2], 'w', newline='')); hdr = None; n = 0
for r in ws.iter_rows(values_only=True):
    if hdr is None:
        if r and r[0] and str(r[0]).strip().upper().startswith('SUBST'): hdr = r; w.writerow(r)
        continue
    w.writerow(r); n += 1
print(n, 'linhas')
